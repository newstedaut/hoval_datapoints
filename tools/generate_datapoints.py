#!/usr/bin/env python3
"""Generates datapoints.csv from Hoval's official TopTronic E Modbus datapoint list (xlsx).

The xlsx is Hoval AG's copyrighted spreadsheet and is intentionally NOT part of
this repository. This project only ships a generator that turns your own local
copy into an open, structured CSV -- so this repo never re-hosts Hoval's file,
and stays in sync whenever Hoval updates it.

Download (current, as of 2026):
  https://cdn.hoval.com/toptronice-gateway-modbus-datapoints_hybris_original.xlsx
  (mirror path sometimes used: https://www.hoval.com/misc/TTE/TTE-GW-Modbus-datapoints.xlsx)

Usage:
  python3 generate_datapoints.py <datapoint-list.xlsx> [unit_ids] [out.csv]

  unit_ids   comma-separated Modbus unit IDs to include, default "1,520,143"
             (1 = WEZ/heat generator, 520 = ventilation HV, 143 = buffer module PS
              -- other bus addresses: see the unit-ID table in Hoval's own manual)

Output columns:
  reg, unit_id, fg, fn, dp, name, type, decimal, min, max, writable, unit

This is the same approach already used in production by
https://github.com/newstedaut/HoxPi (tools/gen_registers.py) -- ported here so
every project consuming this repo's CSV can regenerate it independently instead
of trusting a re-hosted copy.
"""
import csv
import sys


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    xlsx = sys.argv[1]
    unit_ids = {int(x) for x in (sys.argv[2] if len(sys.argv) > 2 else "1,520,143").split(",")}
    out_path = sys.argv[3] if len(sys.argv) > 3 else "datapoints.csv"

    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    sheet = next((s for s in wb.sheetnames if s.lower().startswith("eng")), wb.sheetnames[0])
    ws = wb[sheet]

    hdr, rows, seen = None, [], set()
    for row in ws.iter_rows(values_only=True):
        if hdr is None:
            hdr = [str(h or "").strip().lower() for h in row]
            i = {h: n for n, h in enumerate(hdr)}
            c = {k: i[k] for k in ("register address", "unitname", "unitid", "functiongroup",
                                    "functionnumber", "datapointid", "datapointname",
                                    "typename", "decimal", "min. value", "max. value",
                                    "writable", "unit")}
            continue
        try:
            reg, uid = int(row[c["register address"]]), int(row[c["unitid"]])
        except (TypeError, ValueError):
            continue
        if uid not in unit_ids or reg in seen:
            continue
        seen.add(reg)

        def num(key):
            v = row[c[key]]
            try:
                return int(v)
            except (TypeError, ValueError):
                return None

        rows.append({
            "reg": reg,
            "unit_id": uid,
            "fg": num("functiongroup") or 0,
            "fn": num("functionnumber") or 0,
            "dp": num("datapointid") or 0,
            "name": str(row[c["datapointname"]] or ""),
            "type": str(row[c["typename"]] or ""),
            "decimal": num("decimal") or 0,
            "min": num("min. value"),
            "max": num("max. value"),
            "writable": str(row[c["writable"]] or "No"),
            "unit": str(row[c["unit"]] or ""),
        })
    rows.sort(key=lambda r: r["reg"])

    with open(out_path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["reg", "unit_id", "fg", "fn", "dp", "name",
                                          "type", "decimal", "min", "max", "writable", "unit"])
        w.writeheader()
        w.writerows(rows)

    print(f"{out_path}: {len(rows)} datapoints (unit IDs {sorted(unit_ids)}, sheet '{sheet}')")


if __name__ == "__main__":
    main()
